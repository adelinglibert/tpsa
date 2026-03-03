"""
build_template.py – Generate a professional template.docx with all TPSA variables.
Run once: python3 build_template.py
"""

from docx import Document
from docx.shared import Pt, RGBColor, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

OUTPUT = "template.docx"

# ── Colour palette ──────────────────────────────────────────────────────────
DARK_NAVY   = RGBColor(0x0D, 0x1B, 0x2A)   # near-black navy
MID_BLUE    = RGBColor(0x1B, 0x4F, 0x72)   # section headers
ACCENT_BLUE = RGBColor(0x21, 0x8B, 0xC3)   # accents / borders
LIGHT_GREY  = RGBColor(0xF2, 0xF4, 0xF7)   # table row shading
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
RED_RAG     = RGBColor(0xC0, 0x39, 0x2B)
AMBER_RAG   = RGBColor(0xD0, 0x7A, 0x10)
GREEN_RAG   = RGBColor(0x1E, 0x87, 0x49)

def _shade_cell(cell, hex_fill: str):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement("w:shd")
    shd.set(qn("w:val"),   "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"),  hex_fill)
    tcPr.append(shd)

def _set_cell_border(cell, **kwargs):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        val = kwargs.get(side)
        if val:
            bdr = OxmlElement(f"w:{side}")
            for k, v in val.items():
                bdr.set(qn(k), v)
            tcBorders.append(bdr)
    tcPr.append(tcBorders)

def _set_font(run, name="Calibri", size=10, bold=False, italic=False,
              color: RGBColor = None):
    run.font.name    = name
    run.font.size    = Pt(size)
    run.font.bold    = bold
    run.font.italic  = italic
    if color:
        run.font.color.rgb = color

def _heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if level == 1 else 8)
    p.paragraph_format.space_after  = Pt(4)
    run = p.add_run(text)
    if level == 1:
        _set_font(run, size=15, bold=True, color=MID_BLUE)
        p.paragraph_format.left_indent = Cm(0)
        # bottom border
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"),  "6")
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), "218BC3")
        pBdr.append(bottom)
        pPr.append(pBdr)
    else:
        _set_font(run, size=12, bold=True, color=MID_BLUE)
    return p

def _body(doc, text, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after  = Pt(4)
    run = p.add_run(text)
    _set_font(run, italic=italic)
    return p

def _label_value(doc, label, value_var):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after  = Pt(1)
    r1 = p.add_run(f"{label}: ")
    _set_font(r1, bold=True, size=10)
    r2 = p.add_run(value_var)
    _set_font(r2, size=10)
    return p

def _simple_table(doc, headers, widths=None):
    """Create a styled table with a dark header row."""
    n_cols = len(headers)
    tbl = doc.add_table(rows=1, cols=n_cols)
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl.style = "Table Grid"
    if widths:
        for i, w in enumerate(widths):
            tbl.columns[i].width = Inches(w)
    hdr_row = tbl.rows[0]
    for i, h in enumerate(headers):
        cell = hdr_row.cells[i]
        _shade_cell(cell, "1B4F72")
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        para = cell.paragraphs[0]
        para.paragraph_format.space_before = Pt(3)
        para.paragraph_format.space_after  = Pt(3)
        run = para.add_run(h)
        _set_font(run, bold=True, size=9, color=WHITE)
    return tbl

def _add_row(tbl, values, shade=False):
    row = tbl.add_row()
    for i, val in enumerate(values):
        cell = row.cells[i]
        if shade:
            _shade_cell(cell, "F2F4F7")
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        para = cell.paragraphs[0]
        para.paragraph_format.space_before = Pt(2)
        para.paragraph_format.space_after  = Pt(2)
        run = para.add_run(str(val))
        _set_font(run, size=9)
    return row

def _page_break(doc):
    doc.add_page_break()

# ── Build document ───────────────────────────────────────────────────────────
doc = Document()

# --- Page margins ---
for section in doc.sections:
    section.top_margin    = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin   = Cm(2.5)
    section.right_margin  = Cm(2.5)

# ════════════════════════════════════════════════════════════
# COVER PAGE
# ════════════════════════════════════════════════════════════
title_p = doc.add_paragraph()
title_p.paragraph_format.space_before = Pt(60)
title_p.paragraph_format.space_after  = Pt(6)
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
tr = title_p.add_run("{{ title }}")
_set_font(tr, size=26, bold=True, color=MID_BLUE)

sub_p = doc.add_paragraph()
sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub_p.paragraph_format.space_after = Pt(4)
sr = sub_p.add_run("{{ company_name }}")
_set_font(sr, size=18, bold=False, color=DARK_NAVY)

date_p = doc.add_paragraph()
date_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
date_p.paragraph_format.space_after = Pt(60)
dr = date_p.add_run("Assessment Date: {{ assessment_date }}")
_set_font(dr, size=11, italic=True, color=DARK_NAVY)

intro_p = doc.add_paragraph()
intro_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
ir = intro_p.add_run("{{ introduction }}")
_set_font(ir, size=10, italic=True, color=DARK_NAVY)

_page_break(doc)

# ════════════════════════════════════════════════════════════
# 1. COMPANY OVERVIEW
# ════════════════════════════════════════════════════════════
_heading(doc, "1. Company Overview")
_heading(doc, "1.1 Services", level=2)
_body(doc, "{{ services }}")
_heading(doc, "1.2 Technical Setup & Architecture", level=2)
_body(doc, "{{ technical_setup }}")
_heading(doc, "1.3 Security Measures", level=2)
_body(doc, "{{ security_measures }}")
_heading(doc, "1.4 Data Flows", level=2)
_body(doc, "{{ data_flows }}")
_heading(doc, "1.5 Third Parties & Sub-processors", level=2)
_body(doc, "{{ third_parties }}")

_page_break(doc)

# ════════════════════════════════════════════════════════════
# 2. EXECUTIVE SCORECARD (Feature 3 – conditional)
# ════════════════════════════════════════════════════════════
_heading(doc, "2. Executive Scorecard")
_label_value(doc, "Overall Rating", "{{ scorecard_overall_rating }}")
_body(doc, "{{ scorecard_overall_summary }}")

doc.add_paragraph()
sc_tbl = _simple_table(doc, ["Domain", "Score", "vs. Baseline", "Key Finding", "Source"],
                        widths=[1.5, 0.8, 0.9, 2.6, 1.2])
_body(doc, "{% for d in scorecard_domains %}")
_add_row(sc_tbl, [
    "{{ d.domain }}", "{{ d.score }}", "{{ d.delta }}",
    "{{ d.key_finding }}", "{{ d.evidence_source }}"
])
_body(doc, "{% endfor %}")

_page_break(doc)

# ════════════════════════════════════════════════════════════
# 3. RED FLAGS (Feature B – conditional)
# ════════════════════════════════════════════════════════════
_heading(doc, "3. Red Flags")
_body(doc, "{% if red_flags %}", italic=True)

rf_tbl = _simple_table(doc, ["Flag", "Severity", "Evidence", "Implication"],
                        widths=[1.8, 0.7, 2.2, 2.3])
_body(doc, "{% for f in red_flags %}")
_add_row(rf_tbl, [
    "{{ f.flag }}", "{{ f.severity }}", "{{ f.evidence }}", "{{ f.implication }}"
])
_body(doc, "{% endfor %}")
_body(doc, "{% else %}No critical red flags identified in the provided documentation.{% endif %}",
      italic=True)

_page_break(doc)

# ════════════════════════════════════════════════════════════
# 4. RISK REGISTER
# ════════════════════════════════════════════════════════════
_heading(doc, "4. Risk Register")
_body(doc, "{% for r in risks_and_recommendations %}")

risk_tbl = _simple_table(doc,
    ["Risk Title", "Likelihood", "Impact", "Risk Level"],
    widths=[3.5, 1.1, 1.0, 1.0])
_add_row(risk_tbl, [
    "{{ r.title }}", "{{ r.likelihood }}", "{{ r.impact }}", "{{ r.risk_level }}"
])

_body(doc, "Description: {{ r.description }}")
_body(doc, "Recommendations: {{ r.recommendations }}")
_label_value(doc, "Sources", "{{ r.sources }}")
doc.add_paragraph()

_body(doc, "{% endfor %}")

_page_break(doc)

# ════════════════════════════════════════════════════════════
# 5. REMEDIATION PLAN
# ════════════════════════════════════════════════════════════
_heading(doc, "5. Post-Assessment Remediation Plan")

_heading(doc, "5.1 Priority Actions (0–30 days)", level=2)
pa_tbl = _simple_table(doc, ["Action", "Rationale", "Owner", "Deadline"],
                        widths=[2.6, 2.2, 1.2, 1.0])
_body(doc, "{% for a in priority_actions %}")
_add_row(pa_tbl, ["{{ a.action }}", "{{ a.rationale }}", "{{ a.owner }}", "{{ a.deadline }}"])
_body(doc, "{% endfor %}")

doc.add_paragraph()
_heading(doc, "5.2 Medium-Term Actions (1–6 months)", level=2)
mt_tbl = _simple_table(doc, ["Action", "Rationale", "Owner", "Deadline"],
                        widths=[2.6, 2.2, 1.2, 1.0])
_body(doc, "{% for a in medium_term_actions %}")
_add_row(mt_tbl, ["{{ a.action }}", "{{ a.rationale }}", "{{ a.owner }}", "{{ a.deadline }}"])
_body(doc, "{% endfor %}")

doc.add_paragraph()
_heading(doc, "5.3 Strategic Actions (6–18 months)", level=2)
st_tbl = _simple_table(doc, ["Action", "Rationale", "Owner", "Deadline"],
                        widths=[2.6, 2.2, 1.2, 1.0])
_body(doc, "{% for a in strategic_actions %}")
_add_row(st_tbl, ["{{ a.action }}", "{{ a.rationale }}", "{{ a.owner }}", "{{ a.deadline }}"])
_body(doc, "{% endfor %}")

doc.add_paragraph()
_heading(doc, "5.4 Governance Recommendations", level=2)
_body(doc, "{{ governance_recommendations }}")
_heading(doc, "5.5 Ownership Matrix", level=2)
_body(doc, "{{ ownership_matrix }}")

_page_break(doc)

# ════════════════════════════════════════════════════════════
# 6. EFFORT ESTIMATE (Feature 2 – conditional)
# ════════════════════════════════════════════════════════════
_heading(doc, "6. Remediation Effort Estimate")
_label_value(doc, "Organisation Size", "{{ effort_org_size }}")
_label_value(doc, "Total Effort (man-days)",
             "Min {{ effort_total_min }} / Realistic {{ effort_total_mid }} / Max {{ effort_total_max }}")
_body(doc, "Assumptions: {{ effort_key_assumptions }}", italic=True)

doc.add_paragraph()
eff_tbl = _simple_table(doc,
    ["Action", "Phase", "Min", "Mid", "Max", "Role", "In-House?", "Notes"],
    widths=[2.0, 0.7, 0.4, 0.4, 0.4, 1.2, 0.7, 1.2])
_body(doc, "{% for a in effort_action_estimates %}")
_add_row(eff_tbl, [
    "{{ a.action }}", "{{ a.phase }}",
    "{{ a.effort_days_min }}", "{{ a.effort_days_mid }}", "{{ a.effort_days_max }}",
    "{{ a.role_profile }}", "{{ a.in_house_feasible }}", "{{ a.notes }}"
])
_body(doc, "{% endfor %}")

_page_break(doc)

# ════════════════════════════════════════════════════════════
# 7. IT STACK COMPARISON (Feature 1 – conditional)
# ════════════════════════════════════════════════════════════
_heading(doc, "7. IT Stack Comparison")
_label_value(doc, "Overall Integration Risk", "{{ integration_risk }}")
_body(doc, "{{ integration_summary }}")

doc.add_paragraph()
_heading(doc, "7.1 Supplier Stack", level=2)
stack_tbl = _simple_table(doc, ["Layer", "Supplier Tool"], widths=[2.0, 5.0])
for layer, var in [
    ("IAM",           "{{ supplier_stack_iam }}"),
    ("SIEM",          "{{ supplier_stack_siem }}"),
    ("ITSM",          "{{ supplier_stack_itsm }}"),
    ("Cloud",         "{{ supplier_stack_cloud }}"),
    ("Network",       "{{ supplier_stack_network }}"),
    ("Endpoint",      "{{ supplier_stack_endpoint }}"),
    ("Data Platform", "{{ supplier_stack_data }}"),
    ("BCDR",          "{{ supplier_stack_bcdr }}"),
]:
    _add_row(stack_tbl, [layer, var], shade=(layer in ("IAM","Cloud","Endpoint","BCDR")))

doc.add_paragraph()
_heading(doc, "7.2 Compatibility Matrix", level=2)
compat_tbl = _simple_table(doc,
    ["Layer", "Supplier", "Acquirer", "Status", "Integration Hurdle", "Consolidation Opportunity"],
    widths=[0.9, 1.1, 1.1, 0.9, 1.8, 1.2])
_body(doc, "{% for row in compatibility_matrix %}")
_add_row(compat_tbl, [
    "{{ row.layer }}", "{{ row.supplier_tool }}", "{{ row.acquirer_tool }}",
    "{{ row.status }}", "{{ row.hurdle }}", "{{ row.opportunity }}"
])
_body(doc, "{% endfor %}")

_page_break(doc)

# ════════════════════════════════════════════════════════════
# 8. CONTRACTUAL RISK ANALYSIS (Feature A – conditional)
# ════════════════════════════════════════════════════════════
_heading(doc, "8. Contractual Risk Analysis")
_label_value(doc, "Documents Reviewed", "{{ contractual_docs_reviewed }}")
_label_value(doc, "Contractual Risk Rating", "{{ contractual_risk_rating }}")
_body(doc, "{{ contractual_summary }}")

doc.add_paragraph()
clause_tbl = _simple_table(doc,
    ["Clause", "Golden Standard", "Current Status", "Risk if Absent"],
    widths=[1.6, 2.2, 1.0, 2.2])
_body(doc, "{% for c in missing_clauses %}")
_add_row(clause_tbl, [
    "{{ c.clause_title }}", "{{ c.golden_standard }}",
    "{{ c.current_status }}", "{{ c.risk_if_absent }}"
])
_body(doc, "{% endfor %}")

_page_break(doc)

# ════════════════════════════════════════════════════════════
# 9. SUPPLY CHAIN MAP (Feature C – conditional)
# ════════════════════════════════════════════════════════════
_heading(doc, "9. Supply Chain Map")
_label_value(doc, "Overall Supply Chain Risk", "{{ supply_chain_risk }}")
_body(doc, "{{ supply_chain_summary }}")

doc.add_paragraph()
sc_map_tbl = _simple_table(doc,
    ["Party", "Tier", "Category", "Data Access", "Location", "Concentration Risk"],
    widths=[1.3, 0.7, 1.0, 0.9, 1.0, 2.1])
_body(doc, "{% for p in supply_chain_parties %}")
_add_row(sc_map_tbl, [
    "{{ p.name }}", "{{ p.tier }}", "{{ p.service_category }}",
    "{{ p.data_access }}", "{{ p.location }}", "{{ p.concentration_risk }}"
])
_body(doc, "{% endfor %}")

_page_break(doc)

# ════════════════════════════════════════════════════════════
# 10. DATA SENSITIVITY HEATMAP (Feature E – conditional)
# ════════════════════════════════════════════════════════════
_heading(doc, "10. Data Sensitivity Heatmap")
_label_value(doc, "DPIA Required", "{{ dpia_required }}")
_body(doc, "DPIA Rationale: {{ dpia_rationale }}", italic=True)
_body(doc, "{{ heatmap_summary }}")

doc.add_paragraph()
heat_tbl = _simple_table(doc,
    ["Data Type", "Sensitivity Tier", "Regulation", "Cross-Border", "Destinations", "Retention"],
    widths=[1.6, 1.2, 0.9, 0.8, 1.3, 1.2])
_body(doc, "{% for d in data_types %}")
_add_row(heat_tbl, [
    "{{ d.data_type }}", "{{ d.sensitivity_tier }}", "{{ d.regulation }}",
    "{{ d.cross_border }}", "{{ d.destinations }}", "{{ d.retention }}"
])
_body(doc, "{% endfor %}")

_page_break(doc)

# ════════════════════════════════════════════════════════════
# 11. FOLLOW-UP QUESTIONNAIRE (Feature D – conditional)
# ════════════════════════════════════════════════════════════
_heading(doc, "11. Follow-up Questionnaire")
_body(doc, "The following targeted questions were generated based on the gaps identified above.",
      italic=True)

doc.add_paragraph()
q_tbl = _simple_table(doc,
    ["#", "Question", "Gap Addressed", "Control Ref.", "Expected Evidence"],
    widths=[0.3, 2.2, 1.8, 1.0, 1.7])
_body(doc, "{% for q in followup_questions %}")
_add_row(q_tbl, [
    "{{ loop.index }}",
    "{{ q.question }}", "{{ q.gap_addressed }}",
    "{{ q.control_reference }}", "{{ q.expected_evidence }}"
])
_body(doc, "{% endfor %}")

doc.add_paragraph()
_body(doc, "— End of Memorandum —", italic=True)

doc.save(OUTPUT)
print(f"[OK] Template saved to {OUTPUT}")
